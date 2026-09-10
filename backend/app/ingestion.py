from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import fitz
from docx import Document as DocxDocument

from app.config import settings
from app.knowledge import get_base
from app.store import vector_store

DOC_DIR = Path("data/documents")
DOC_DIR.mkdir(parents=True, exist_ok=True)
SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}
MARKDOWN_HEADING = re.compile(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$")


def chunk_text(text: str, size: int | None = None, overlap: int | None = None) -> list[str]:
    size = size or settings.chunk_size
    overlap = overlap or settings.chunk_overlap
    clean = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not clean:
        return []

    chunks: list[str] = []
    start = 0
    while start < len(clean):
        end = min(len(clean), start + size)
        if end < len(clean):
            candidates = [
                clean.rfind(mark, start + size // 2, end)
                for mark in ("\n", "。", "；", "！", "？", ". ")
            ]
            cut = max(candidates)
            if cut > start:
                end = cut + 1

        value = clean[start:end].strip()
        if value:
            chunks.append(value)
        if end >= len(clean):
            break
        start = max(start + 1, end - overlap)
    return chunks


def chunk_markdown(text: str) -> list[dict[str, Any]]:
    """Split Markdown by heading hierarchy, then chunk body-bearing sections.

    Parent headings are carried into child-section metadata and Citation content.
    Adjacent headings do not emit empty heading-only chunks.
    """
    sections: list[tuple[str, list[str]]] = []
    heading_stack: list[tuple[int, str, str]] = []
    current_lines: list[str] = []

    def flush() -> None:
        nonlocal current_lines
        if not any(line.strip() for line in current_lines):
            current_lines = []
            return
        section_title = " > ".join(title for _, title, _ in heading_stack)
        heading_lines = [raw for _, _, raw in heading_stack]
        sections.append((section_title, [*heading_lines, *current_lines]))
        current_lines = []

    for raw_line in text.splitlines():
        match = MARKDOWN_HEADING.match(raw_line)
        if match:
            flush()
            level = len(match.group(1))
            title = match.group(2).strip()
            while heading_stack and heading_stack[-1][0] >= level:
                heading_stack.pop()
            heading_stack.append((level, title, raw_line.strip()))
        else:
            current_lines.append(raw_line)
    flush()

    if not sections:
        return [
            {"content": chunk, "page": None, "chunk_index": index, "section_title": None}
            for index, chunk in enumerate(chunk_text(text))
        ]

    parts: list[dict[str, Any]] = []
    index = 0
    for heading, lines in sections:
        section_text = "\n".join(lines)
        for chunk in chunk_text(section_text):
            parts.append(
                {
                    "content": chunk,
                    "page": None,
                    "chunk_index": index,
                    "section_title": heading or None,
                }
            )
            index += 1
    return parts


def parse_document(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError("仅支持 PDF / DOCX / TXT / Markdown")

    parts: list[dict[str, Any]] = []

    if suffix == ".pdf":
        pdf = fitz.open(path)
        try:
            for page_no, page in enumerate(pdf, start=1):
                text = page.get_text("text")
                for index, chunk in enumerate(chunk_text(text)):
                    parts.append(
                        {
                            "content": chunk,
                            "page": page_no,
                            "chunk_index": index,
                            "section_title": None,
                        }
                    )
        finally:
            pdf.close()
        return parts

    if suffix == ".docx":
        doc = DocxDocument(path)
        blocks: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(blocks)
    else:
        text = path.read_text(encoding="utf-8", errors="ignore")

    if suffix == ".md":
        return chunk_markdown(text)

    return [
        {"content": chunk, "page": None, "chunk_index": index, "section_title": None}
        for index, chunk in enumerate(chunk_text(text))
    ]


def document_path(knowledge_base_id: str, file_name: str) -> Path:
    directory = DOC_DIR / knowledge_base_id
    directory.mkdir(parents=True, exist_ok=True)
    return directory / Path(file_name).name


def ingest_file(path: Path, knowledge_base_id: str) -> int:
    base = get_base(knowledge_base_id)
    chunks = parse_document(path)
    if not chunks:
        raise ValueError("文档未解析出有效文本")

    vector_store.delete_file(path.name, knowledge_base_id)
    enriched = [
        {
            **chunk,
            "file_name": path.name,
            "file_type": path.suffix.lower().lstrip("."),
            "document_title": path.stem,
            "knowledge_base_id": knowledge_base_id,
            "knowledge_base_name": base["name"],
            "retrieval_schema_version": settings.retrieval_schema_version,
        }
        for chunk in chunks
    ]
    return vector_store.add_chunks(enriched)
